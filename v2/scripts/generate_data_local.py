#!/usr/bin/env python3
"""
generate_data_local.py
======================
Genera data/data.json leyendo los Excel locales.
Ejecutar con doble clic o: python scripts/generate_data_local.py

No requiere Azure ni SharePoint — usa las rutas locales.
"""

import os, sys, json, glob, shutil, tempfile, re
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    print("Instalando openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '--break-system-packages', '-q'])
    import openpyxl

# win32com para leer Outlook (solo Windows con Outlook instalado)
try:
    import win32com.client as _w32
    WIN32_OK = True
except ImportError:
    WIN32_OK = False

# ─── Rutas ──────────────────────────────────────────────────────────────────
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # v2/

EXCEL_INCIDENCIAS = os.path.join(
    os.path.expanduser('~'), 'Software Broker',
    'Nathalia Moreno - Ordenes incidentadas 2026',
    'Ordenes con novedad Junio - dic.xlsx'
)
EXCEL_INGRESADAS_DIR = os.path.join(
    os.path.expanduser('~'), 'Software Broker',
    'Nathalia Moreno - Ordenes incidentadas 2026',
    'Ingresadas diarias'
)
OUTPUT_PATH = os.path.join(BASE, 'data', 'data.json')


# ─── Helpers ────────────────────────────────────────────────────────────────
def fmt_date(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    # intentar parsear si viene como string
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(s[:10], fmt[:10]).strftime('%Y-%m-%d')
        except Exception:
            pass
    return s[:10]

def clean(v):
    if v is None:
        return ''
    return str(v).strip()

def extract_pais(marca_pais, marca):
    """Extrae el país de 'Marca PAIS' quitando la marca."""
    mp = clean(marca_pais)
    m  = clean(marca)
    if mp.upper().startswith(m.upper()):
        pais = mp[len(m):].strip()
    else:
        pais = mp
    return pais or mp


# ─── Leer incidencias ────────────────────────────────────────────────────────
def open_workbook(path):
    """Abre el Excel copiándolo a un temp para evitar PermissionError si está abierto."""
    tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    tmp.close()
    try:
        shutil.copy2(path, tmp.name)
        return openpyxl.load_workbook(tmp.name, read_only=True, data_only=True)
    except Exception:
        os.unlink(tmp.name)
        raise

def read_incidencias():
    print(f'\nLeyendo: {EXCEL_INCIDENCIAS}')
    if not os.path.exists(EXCEL_INCIDENCIAS):
        print(f'  [ERROR] No existe: {EXCEL_INCIDENCIAS}')
        return []

    wb = open_workbook(EXCEL_INCIDENCIAS)
    ws = wb.active
    records = []
    headers = None

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = list(row)
            continue
        if all(v is None for v in row[:6]):
            break
        rec = dict(zip(headers, row))

        marca     = clean(rec.get('marca'))
        marca_pais= clean(rec.get('Marca PAIS'))
        pais      = extract_pais(marca_pais, marca)
        fecha     = fmt_date(rec.get('fecha de orden'))
        # estado Y com = 'Comentario continuidad'
        com_cont  = clean(rec.get('Comentario continuidad'))
        ops       = clean(rec.get('Duplicado'))
        ov        = clean(rec.get('VENTA'))
        ticket    = clean(rec.get('Ticket'))
        pendiente = clean(rec.get('Pendiente por:'))
        proveedor = clean(rec.get('Proveedor'))
        canal     = clean(rec.get('Canal '))

        if not fecha and not ov:
            continue

        records.append({
            'fecha':     fecha,
            'fn':        fmt_date(rec.get('Fecha de notificacion JIRA')),  # fecha notif JIRA
            'marca':     marca,
            'pais':      pais,
            'estado':    com_cont,   # Comentario continuidad
            'com':       com_cont,   # mismo — para filtros isSinStock / isKrono
            'ops':       ops,        # Duplicado / Acción OPS (para isEnRevIT)
            'pen':       clean(rec.get('Pendiente por:')),  # Pendiente por
            'ov':        ov,
            'ticket':    ticket,
            'pendiente': pendiente,
            'proveedor': proveedor,
            'canal':     canal,
            'detalle':   clean(rec.get('Estado Caso')),
        })

    wb.close()
    print(f'  {len(records)} registros de incidencias')
    return records


# ─── Leer IX desde Outlook ──────────────────────────────────────────────────
SUBJECT_PREFIX = 'Reporte Diario de Ordenes Drivers'   # sin tildes para matching robusto

def _norm_subject(s):
    """Quita tildes y pasa a minúsculas para comparar asuntos."""
    import unicodedata
    return ''.join(
        c for c in unicodedata.normalize('NFD', s.lower())
        if unicodedata.category(c) != 'Mn'
    )

def read_ix_from_outlook():
    """
    Abre Outlook, busca correos 'Reporte Diario de Ordenes Drivers',
    lee el Excel adjunto y devuelve:
      {
        'porDia':  {'2026-09-01': 1114, '2026-09-02': 980, ...},
        'porMes':  {'2026-09': 5432, ...}
      }
    Columnas esperadas en el Excel: Fecha | Driver 1 | Driver 2 | CIXC | Total
    """
    if not WIN32_OK:
        print('\n  [INFO] pywin32 no disponible — omitiendo lectura de Outlook.')
        print('         Instala con: pip install pywin32')
        return {'porDia': {}, 'porMes': {}}

    print('\nLeyendo correos Outlook (Reporte Diario de Ordenes Drivers)...')

    ixc_por_dia = {}   # 'YYYY-MM-DD' -> int total

    try:
        outlook = _w32.Dispatch('Outlook.Application')
        ns      = outlook.GetNamespace('MAPI')
        inbox   = ns.GetDefaultFolder(6)   # 6 = Inbox

        # Buscar carpeta DRIVER recorriendo cuentas con manejo de errores
        carpeta = None
        for store in ns.Folders:
            try:
                for f in store.Folders:
                    try:
                        if f.Name.upper() == 'DRIVER':
                            carpeta = f
                            break
                        for sf in f.Folders:
                            try:
                                if sf.Name.upper() == 'DRIVER':
                                    carpeta = sf
                                    break
                            except Exception:
                                pass
                    except Exception:
                        pass
                    if carpeta:
                        break
            except Exception:
                pass
            if carpeta:
                break

        if carpeta is None:
            print('  [WARN] Carpeta DRIVER no encontrada — buscando en Bandeja de entrada')
            carpeta = inbox

        # Obtener items — sin Restrict para evitar errores de sintaxis COM
        items = carpeta.Items
        try:
            items.Sort('[ReceivedTime]', True)   # más reciente primero
        except Exception:
            pass

        prefix_norm = _norm_subject(SUBJECT_PREFIX)
        procesados  = 0

        for mail in items:
            try:
                subj_norm = _norm_subject(mail.Subject or '')
            except Exception:
                continue
            if prefix_norm not in subj_norm:
                continue

            for att in mail.Attachments:
                fname = att.FileName or ''
                if not fname.lower().endswith('.xlsx'):
                    continue

                tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
                tmp.close()
                try:
                    att.SaveAsFile(tmp.name)
                    wb = openpyxl.load_workbook(tmp.name, read_only=True, data_only=True)
                    ws = wb.active
                    headers = None
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        if i == 0:
                            headers = [str(h).strip() if h else '' for h in row]
                            continue
                        if not headers or all(v is None for v in row):
                            break
                        rec = dict(zip(headers, row))
                        fecha = fmt_date(rec.get('Fecha'))
                        if not fecha or not re.match(r'^\d{4}-\d{2}-\d{2}$', fecha):
                            continue   # fila de totales u otra sin fecha válida
                        total = rec.get('Total')
                        if total is None:
                            continue
                        try:
                            ixc_por_dia[fecha] = int(float(total))
                        except (ValueError, TypeError):
                            pass
                    wb.close()
                    procesados += 1
                except Exception as e:
                    print(f'    [WARN] No se pudo leer adjunto "{fname}": {e}')
                finally:
                    try:
                        os.unlink(tmp.name)
                    except Exception:
                        pass
                break   # solo el primer .xlsx del correo

        print(f'  {procesados} correo(s) procesados — {len(ixc_por_dia)} dias con dato IX')

    except Exception as e:
        print(f'  [WARN] Error al conectar con Outlook: {e}')
        return {'porDia': {}, 'porMes': {}}

    # Calcular acumulados por mes
    ixc_por_mes = {}
    for fecha, total in ixc_por_dia.items():
        mes = fecha[:7]   # 'YYYY-MM'
        ixc_por_mes[mes] = ixc_por_mes.get(mes, 0) + total

    return {'porDia': ixc_por_dia, 'porMes': ixc_por_mes}


# ─── Leer ingresadas diarias ─────────────────────────────────────────────────
def read_ingresadas():
    pattern = os.path.join(EXCEL_INGRESADAS_DIR, '*.xlsx')
    files   = glob.glob(pattern)
    print(f'\nBuscando ingresadas en: {EXCEL_INGRESADAS_DIR}')
    print(f'  {len(files)} archivos encontrados')

    records = []
    for fpath in sorted(files):
        fname = os.path.basename(fpath)
        try:
            wb = open_workbook(fpath)
            ws = wb.active
            headers = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = list(row)
                    continue
                if all(v is None for v in row):
                    break
                if headers is None:
                    continue
                rec = dict(zip(headers, row))

                # Intentar múltiples nombres de columna para la fecha
                fecha = (fmt_date(rec.get('fecha de orden'))
                      or fmt_date(rec.get('Fecha de orden'))
                      or fmt_date(rec.get('fecha notificacion jira'))
                      or '')

                ov    = clean(rec.get('VENTA') or rec.get('venta') or '')
                pais  = clean(rec.get('PAIS') or rec.get('Pais') or rec.get('pais') or '')
                marca = clean(rec.get('marca') or rec.get('Marca') or '')

                if not fecha:
                    continue

                records.append({
                    'fecha': fecha,
                    'ov':    ov,
                    'pais':  pais,
                    'marca': marca,
                })
            wb.close()
        except Exception as e:
            print(f'  [WARN] {fname}: {e}')

    print(f'  {len(records)} registros de ingresadas')
    return records


# ─── Main ────────────────────────────────────────────────────────────────────
def main():
    print('=' * 50)
    print('generate_data_local.py')
    print(f'Inicio: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    print('=' * 50)

    incidencias = read_incidencias()
    ingresadas  = read_ingresadas()
    ixc         = read_ix_from_outlook()

    data = {
        'generatedAt': datetime.now(timezone.utc).isoformat(),
        'source': 'local',
        'incidencias': incidencias,
        'ingresadas':  ingresadas,
        'jiras':       [],
        'ixcData':     ixc['porDia'],   # {'2026-09-01': 1114, ...}
        'ixcMes':      ixc['porMes'],   # {'2026-09': 5432, ...}
    }

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)

    # Escribir data.json (para servidor HTTP / Azure)
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    # Escribir data.js (para abrir index.html directamente sin servidor)
    js_path = OUTPUT_PATH.replace('.json', '.js')
    with open(js_path, 'w', encoding='utf-8') as f:
        f.write('// Generado automáticamente — no editar\n')
        f.write('window.DASHBOARD_DATA = ')
        json.dump(data, f, ensure_ascii=False)
        f.write(';\n')

    size_kb = os.path.getsize(OUTPUT_PATH) / 1024
    n_ixc   = len(ixc['porDia'])
    print(f'\n✓ Archivos generados:')
    print(f'  • {len(incidencias)} incidencias | {len(ingresadas)} ingresadas | {n_ixc} dias IX desde Outlook')
    print(f'  • data.json  → {size_kb:.1f} KB')
    print(f'  • data.js    → {os.path.getsize(js_path)/1024:.1f} KB')
    if n_ixc == 0:
        print(f'\n  [!] No se encontraron datos IX. Verifica que Outlook este abierto')
        print(f'      y que los correos "Reporte Diario de Ordenes Drivers" esten en Bandeja de entrada.')
    print(f'\nAbre src/index.html directamente en el navegador.')

if __name__ == '__main__':
    main()
